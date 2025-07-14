# Start of FINAL corrected code for lecturer_portal/views.py

from django.http import JsonResponse, HttpResponse, HttpResponseNotAllowed, HttpResponseForbidden
from django.contrib.auth import login, logout
from django.contrib.auth.decorators import login_required
from django.views.decorators.csrf import csrf_exempt
from django.utils import timezone
from django.db import transaction
from django.shortcuts import render
import json
import re
import traceback
import io

from admin_portal.models import LecturerCode
from .models import Exam, Question, Option
from student_portal.models import ExamAttempt

# Direct imports - if these fail, the server will crash with a clear ImportError
from docx import Document
import openpyxl

def is_lecturer_authenticated(request):
    """Helper function to check if a user is an authenticated lecturer."""
    if request.user.is_authenticated:
        if not request.user.is_staff and not request.user.is_superuser:
            return True
    return False


@csrf_exempt
def lecturer_code_login_api(request):
    """API for lecturers to log in using a generated code."""
    if request.method != 'POST':
        return HttpResponseNotAllowed(['POST'])
    try:
        data = json.loads(request.body)
        code_value = data.get('code')
    except json.JSONDecodeError:
        return JsonResponse({'error': 'Invalid JSON format.'}, status=400)

    if not code_value:
        return JsonResponse({'error': 'Code is required.'}, status=400)

    try:
        lecturer_code_obj = LecturerCode.objects.select_related('lecturer_user').get(code=code_value, is_active=True)
    except LecturerCode.DoesNotExist:
        return JsonResponse({'error': 'Invalid or inactive code.'}, status=403)

    if lecturer_code_obj.expires_at and lecturer_code_obj.expires_at < timezone.now():
        lecturer_code_obj.is_active = False
        lecturer_code_obj.save()
        return JsonResponse({'error': 'Code has expired.'}, status=403)

    if not lecturer_code_obj.lecturer_user or not lecturer_code_obj.lecturer_user.is_active:
        return JsonResponse({'error': 'Associated lecturer account is not properly configured or inactive.'}, status=403)

    login(request, lecturer_code_obj.lecturer_user)
    return JsonResponse({
        'message': 'Login successful.',
        'username': lecturer_code_obj.lecturer_user.username,
        'user_id': lecturer_code_obj.lecturer_user.id
    })


@csrf_exempt
def lecturer_logout_api(request):
    """API for lecturers to log out."""
    if request.method != 'POST':
        return HttpResponseNotAllowed(['POST'])
    logout(request)
    return JsonResponse({'message': 'Logout successful.'})


def _save_parsed_question(exam_obj, q_text, q_type, options_list, correct_letter, lecturer_answer_key_text):
    """Helper function to save a question and its options to the database."""
    if not q_text:
        print("DEBUG_SAVE: Question text is empty, not saving.")
        return

    final_answer_key = lecturer_answer_key_text if lecturer_answer_key_text is not None else ""
    
    print(f"DEBUG_SAVE: Creating Question: text='{q_text[:30]}...', type='{q_type}', answer_key='{final_answer_key}'")
    question = Question.objects.create(
        exam=exam_obj,
        question_text=q_text,
        question_type=q_type,
        lecturer_answer_key=final_answer_key
    )

    if q_type == 'multiple_choice' and options_list:
        print(f"DEBUG_SAVE: MCQ options_list: {options_list}, correct_letter: {correct_letter}")
        for opt in options_list:
            is_correct = (opt['letter'] == correct_letter)
            Option.objects.create(question=question, option_text=opt['text'], is_correct=is_correct)
            print(f"DEBUG_SAVE:  - Creating Option: text='{opt['text']}', is_correct={is_correct}")
    print(f"DEBUG_SAVE: Question {question.id} saved with type {q_type}.")


@csrf_exempt
def upload_questions_api(request):
    """API for uploading a Word document and parsing it into questions."""
    if not is_lecturer_authenticated(request):
        return JsonResponse({'error': 'Authentication required.'}, status=401)

    if request.method != 'POST':
        return HttpResponseNotAllowed(['POST'])

    exam_id = request.POST.get('exam_id')
    word_file = request.FILES.get('word_file')

    if not all([exam_id, word_file]):
        return JsonResponse({'error': 'exam_id and word_file are required.'}, status=400)

    try:
        exam = Exam.objects.get(id=exam_id, lecturer_user=request.user)
    except Exam.DoesNotExist:
        return JsonResponse({'error': 'Exam not found or you do not have permission.'}, status=404)

    print("--- Starting Question Parsing (Version: FINAL CORRECTED) ---")
    try:
        document = Document(word_file)
        parsed_questions = 0
        with transaction.atomic():
            exam.questions.all().delete()

            current_q_text, current_options, current_ans_letter, current_ans_key = None, [], None, None
            active_q_type = 'multiple_choice'
            header_q_type = None

            for idx, para in enumerate(document.paragraphs):
                text = para.text.strip()
                print(f"DEBUG [{idx}]: Processing: '{text}' || Active Type: {active_q_type} || Header Type: {header_q_type}")
                if not text: continue

                if text.lower() == "essay question":
                    if current_q_text: _save_parsed_question(exam, current_q_text, active_q_type, current_options, current_ans_letter, current_ans_key); parsed_questions += 1
                    current_q_text, current_options, current_ans_letter, current_ans_key = None, [], None, None
                    header_q_type = 'essay'
                    print(f"DEBUG: Header found. Next questions will be type: '{header_q_type}'")
                    continue
                elif text.lower() == "fill in the blanks":
                    if current_q_text: _save_parsed_question(exam, current_q_text, active_q_type, current_options, current_ans_letter, current_ans_key); parsed_questions += 1
                    current_q_text, current_options, current_ans_letter, current_ans_key = None, [], None, None
                    header_q_type = 'fill_in_the_blanks'
                    print(f"DEBUG: Header found. Next questions will be type: '{header_q_type}'")
                    continue

                answer_mcq_match = re.match(r"^Answer:\s*([A-Za-z])$", text, re.IGNORECASE)
                answer_text_match = re.match(r"^Answer:\s*(.+)", text, re.IGNORECASE)

                if current_q_text:
                    if active_q_type == 'multiple_choice' and answer_mcq_match:
                        current_ans_letter = answer_mcq_match.group(1).upper()
                        print(f"DEBUG: Matched Answer for MCQ: '{current_ans_letter}'")
                        continue
                    elif active_q_type != 'multiple_choice' and answer_text_match:
                        current_ans_key = answer_text_match.group(1).strip()
                        print(f"DEBUG: Matched Answer for Non-MCQ: '{current_ans_key}'")
                        _save_parsed_question(exam, current_q_text, active_q_type, [], None, current_ans_key)
                        parsed_questions += 1
                        current_q_text, current_options, current_ans_letter, current_ans_key = None, [], None, None
                        if not header_q_type: active_q_type = 'multiple_choice'
                        else: active_q_type = header_q_type
                        continue

                q_match = re.match(r"^(\d+)\.\s*(.*)", text)
                if q_match:
                    if current_q_text: 
                        _save_parsed_question(exam, current_q_text, active_q_type, current_options, current_ans_letter, current_ans_key); parsed_questions += 1
                    
                    current_options, current_ans_letter, current_ans_key = [], None, None
                    potential_q_text = q_match.group(2).strip()

                    if header_q_type:
                        active_q_type = header_q_type
                        current_q_text = potential_q_text
                        header_q_type = None
                    else:
                        if "essay question:" in potential_q_text.lower():
                            active_q_type = 'essay'
                            current_q_text = re.split(r"essay question:", potential_q_text, flags=re.IGNORECASE, maxsplit=1)[1].strip()
                        elif "fill in the blank:" in potential_q_text.lower() or "____" in potential_q_text:
                            active_q_type = 'fill_in_the_blanks'
                            if "fill in the blank:" in potential_q_text.lower():
                                current_q_text = re.split(r"fill in the blank:", potential_q_text, flags=re.IGNORECASE, maxsplit=1)[1].strip()
                            else:
                                current_q_text = potential_q_text
                        else:
                            active_q_type = 'multiple_choice'
                            current_q_text = potential_q_text
                    print(f"DEBUG: Matched Q (now type '{active_q_type}'): '{current_q_text}'")
                    continue

                option_match = re.match(r"^([A-Za-z])\.\s+(.*)", text)
                if current_q_text and active_q_type == 'multiple_choice' and option_match:
                    current_options.append({'letter': option_match.group(1).upper(), 'text': option_match.group(2).strip()})
                    print(f"DEBUG: Matched OPTION. Options are now: {current_options}")
                    continue
                
                print(f"WARN: Unhandled paragraph: '{text}'")

            if current_q_text:
                _save_parsed_question(exam, current_q_text, active_q_type, current_options, current_ans_letter, current_ans_key); parsed_questions += 1

    except Exception as e:
        print(f"ERROR during parsing: {str(e)} traceback: {traceback.format_exc()}")
        return JsonResponse({'error': f'An error occurred: {str(e)}'}, status=500)

    return JsonResponse({'message': f'Successfully parsed and saved {parsed_questions} questions.'}, status=201)


@csrf_exempt
def update_exam_details_api(request, exam_id):
    if not is_lecturer_authenticated(request):
        return JsonResponse({'error': 'Authentication required.'}, status=401)
    try:
        exam = Exam.objects.get(id=exam_id, lecturer_user=request.user)
    except Exam.DoesNotExist:
        return JsonResponse({'error': 'Exam not found or you do not have permission to modify it.'}, status=404)
    if request.method == 'POST':
        try:
            data = json.loads(request.body)
        except json.JSONDecodeError:
            return JsonResponse({'error': 'Invalid JSON format.'}, status=400)
        updated_fields_count = 0
        if 'title' in data:
            exam.title = data['title']
            updated_fields_count += 1
        if 'course_code' in data:
            exam.course_code = data['course_code']
            updated_fields_count += 1
        if 'start_time' in data:
            try:
                exam.start_time = timezone.datetime.fromisoformat(data['start_time'].replace('Z', '+00:00'))
                updated_fields_count += 1
            except (ValueError, TypeError):
                return JsonResponse({'error': 'Invalid start_time format. Use ISO 8601.'}, status=400)
        if 'end_time' in data:
            try:
                exam.end_time = timezone.datetime.fromisoformat(data['end_time'].replace('Z', '+00:00'))
                updated_fields_count += 1
            except (ValueError, TypeError):
                return JsonResponse({'error': 'Invalid end_time format. Use ISO 8601.'}, status=400)
        if 'duration_minutes' in data:
            try:
                duration = int(data['duration_minutes'])
                if duration <= 0:
                    return JsonResponse({'error': 'Duration must be positive.'}, status=400)
                exam.duration_minutes = duration
                updated_fields_count += 1
            except (ValueError, TypeError):
                return JsonResponse({'error': 'Invalid duration_minutes format. Must be an integer.'}, status=400)
        if 'is_active' in data:
            if isinstance(data['is_active'], bool):
                exam.is_active = data['is_active']
                updated_fields_count += 1
            else:
                return JsonResponse({'error': 'is_active must be a boolean (true/false).'}, status=400)
        if updated_fields_count > 0:
            if exam.start_time and exam.end_time and exam.end_time <= exam.start_time:
                return JsonResponse({'error': 'End time must be after start time.'}, status=400)
            exam.save()
            return JsonResponse({'message': 'Exam details updated successfully.'})
        else:
            return JsonResponse({'message': 'No update data provided or no changes made.'}, status=200)
    return HttpResponseNotAllowed(['POST'])


@csrf_exempt
def get_exam_attempts_api(request, exam_id):
    if not is_lecturer_authenticated(request):
        return JsonResponse({'error': 'Authentication required.'}, status=401)
    try:
        exam = Exam.objects.get(id=exam_id, lecturer_user=request.user)
    except Exam.DoesNotExist:
        return JsonResponse({'error': 'Exam not found or you do not have permission to view it.'}, status=404)
    if request.method == 'GET':
        attempts = ExamAttempt.objects.filter(exam=exam).select_related('student').order_by('-start_time')
        attempts_data = []
        for attempt in attempts:
            student_data = {'name': attempt.student.name, 'matric_number': attempt.student.matric_number}
            attempts_data.append({
                'attempt_id': attempt.id,
                'student': student_data,
                'start_time': attempt.start_time.isoformat() if attempt.start_time else None,
                'submission_time': attempt.submission_time.isoformat() if attempt.submission_time else None,
                'score': attempt.score,
                'grace_login_granted': attempt.grace_login_granted,
                'attempt_deadline': attempt.attempt_deadline.isoformat() if attempt.attempt_deadline else None,
            })
        return JsonResponse({'exam_title': exam.title, 'attempts': attempts_data})
    return HttpResponseNotAllowed(['GET'])


@csrf_exempt
def grant_grace_login_api(request, attempt_id):
    if not is_lecturer_authenticated(request):
        return JsonResponse({'error': 'Authentication required.'}, status=401)
    if request.method == 'POST':
        try:
            exam_attempt = ExamAttempt.objects.select_related('exam').get(id=attempt_id, exam__lecturer_user=request.user)
        except ExamAttempt.DoesNotExist:
            return JsonResponse({'error': 'Exam attempt not found or you do not have permission to modify it.'}, status=404)
        exam_attempt.grace_login_granted = True
        exam_attempt.save(update_fields=['grace_login_granted'])
        return JsonResponse({
            'message': 'Grace login has been successfully granted.',
            'attempt_id': exam_attempt.id,
            'grace_login_granted': exam_attempt.grace_login_granted
        })
    return HttpResponseNotAllowed(['POST'])


@csrf_exempt
def download_results_excel_api(request, exam_id):
    # NOTE: To test this easily, you can temporarily comment out the auth check
    # and the lecturer_user filter in the Exam query as you did before.
    if not is_lecturer_authenticated(request):
        return JsonResponse({'error': 'Authentication required.'}, status=401)
    try:
        exam = Exam.objects.get(id=exam_id, lecturer_user=request.user)
    except Exam.DoesNotExist:
        return JsonResponse({'error': 'Exam not found.'}, status=404)

    if request.method == 'GET':
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"{exam.course_code} Results"

        headers = ["Matric Number", "Student Name", "Start Time", "Submission Time", "Final Score"]
        ws.append(headers)

        attempts = ExamAttempt.objects.filter(exam=exam).select_related('student').order_by('student__matric_number')

        for attempt in attempts:
            row = [
                attempt.student.matric_number,
                attempt.student.name,
                attempt.start_time.strftime("%Y-%m-%d %H:%M:%S") if attempt.start_time else "N/A",
                attempt.submission_time.strftime("%Y-%m-%d %H:%M:%S") if attempt.submission_time else "N/A",
                attempt.score if attempt.score is not None else "Pending"
            ]
            ws.append(row)

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        response = HttpResponse(
            buffer.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{exam.title}_results.xlsx"'
        return response

    return HttpResponseNotAllowed(['GET'])


# --- Views for Frontend Pages --- #

def lecturer_login_page(request):
    """Renders the lecturer login page."""
    return render(request, 'lecturer_portal/login.html')


@login_required
def lecturer_dashboard_page(request):
    """Renders the main lecturer dashboard after login."""
    if not is_lecturer_authenticated(request):
        return HttpResponseForbidden("You do not have access to the lecturer dashboard.")
    
    lecturer_exams = Exam.objects.filter(lecturer_user=request.user).order_by('-created_at')
    
    context = {
        'exams': lecturer_exams,
        'lecturer_name': request.user.first_name or request.user.username
    }
    return render(request, 'lecturer_portal/dashboard.html', context)

# End of corrected code for lecturer_portal/views.py